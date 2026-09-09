"""Plain-text extraction from an uploaded PDF / DOCX / TXT / RTF file.

Image-only (scanned) PDFs are not supported - they need OCR, which needs
system packages. Such a file fails with a clear message rather than
silently producing an empty resume.
"""

from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".txt", ".rtf", ".md"}


class UnreadableFile(Exception):
    """The file cannot be turned into text - message is shown to the user."""


def extract_text(path: str | Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix not in SUPPORTED:
        raise UnreadableFile(
            f"Unsupported file type '{suffix or path.name}'. Supported: {', '.join(sorted(SUPPORTED))}."
        )

    try:
        text = _READERS[suffix](path)
    except UnreadableFile:
        raise
    except Exception as exc:
        raise UnreadableFile(f"Could not read {path.name}: {exc}") from exc

    if not text.strip():
        raise UnreadableFile(
            f"No text found in {path.name}. If it is a scanned or image-only "
            f"document, convert it to a text-based PDF first."
        )
    return text.strip()


def _read_pdf(path: Path) -> str:
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def _read_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    # Plenty of resumes lay everything out inside tables.
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _read_rtf(path: Path) -> str:
    from striprtf.striprtf import rtf_to_text

    return rtf_to_text(path.read_text(encoding="utf-8", errors="ignore"))


def _read_plain(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


_READERS = {
    ".pdf": _read_pdf,
    ".docx": _read_docx,
    ".rtf": _read_rtf,
    ".txt": _read_plain,
    ".md": _read_plain,
}
